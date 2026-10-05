from pathlib import Path
import openpyxl
p = Path(r'd:\Downloads\RASDUMP.xlsx')
print('exists:', p.exists(), 'size:', p.stat().st_size if p.exists() else 'missing')
if p.exists():
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True, max_row=5))
    print('SHEETS:', wb.sheetnames)
    for i, r in enumerate(rows, 1):
        print(f'ROW {i}: {r}')
